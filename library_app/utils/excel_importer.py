"""
Excel & Document importer — reads student, staff, and book data from existing SRM files
"""

import os
import zipfile
import xml.etree.ElementTree as ET
import openpyxl
from pathlib import Path


def _clean(val):
    if val is None:
        return ""
    if isinstance(val, float) and val.is_integer():
        return str(int(val))
    return str(val).strip()


def import_books_from_excel(filepath: str) -> tuple[list[dict], list[str]]:
    """
    Import books from the DEPT LIBRARY FINAL Excel file.
    Returns (list_of_books, list_of_errors)
    """
    books = []
    errors = []

    try:
        wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
    except Exception as e:
        return [], [f"Could not open file: {e}"]

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        if not hasattr(ws, 'max_row'):
            continue  # skip chart sheets

        header = None
        for row in ws.iter_rows(values_only=True):
            if header is None:
                if row and any(str(c).strip().lower() in ["book name", "s.no", "account number"]
                               for c in row if c):
                    header = [str(c).strip() if c else "" for c in row]
                continue

            if not any(row):
                continue

            row_dict = dict(zip(header, row))

            title = _clean(row_dict.get("Book Name", "") or row_dict.get("BOOK NAME", ""))
            if not title or title.lower() in ["book name", "nil", ""]:
                continue

            acc_num = _clean(row_dict.get("Account Number", "") or row_dict.get("ACCOUNT NUMBER", ""))
            publisher = _clean(row_dict.get("Publisher", "") or row_dict.get("PUBLISHER", ""))
            edition = _clean(row_dict.get("Edition", "") or row_dict.get("EDITION", ""))

            # Collect authors
            author_cols = [k for k in header if "author" in k.lower()]
            authors_list = [_clean(row_dict.get(k, "")) for k in author_cols]
            authors = ", ".join(a for a in authors_list if a and a.lower() not in ["nil", "n/a", ""])

            barcode = f"BK{acc_num}" if acc_num else None
            if not barcode:
                continue

            books.append({
                "barcode": barcode,
                "account_number": acc_num,
                "title": title,
                "publisher": publisher,
                "authors": authors,
                "edition": edition,
            })

    wb.close()
    return books, errors


def _import_students_from_docx(filepath: str, year: str, section: str) -> tuple[list[dict], list[str]]:
    patrons = []
    try:
        with zipfile.ZipFile(filepath) as z:
            xml_content = z.read('word/document.xml')
        tree = ET.fromstring(xml_content)
        tables = tree.findall('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tbl')
        for t in tables:
            rows = t.findall('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tr')
            for r in rows:
                cells = r.findall('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tc')
                row_texts = []
                for c in cells:
                    t_nodes = c.findall('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t')
                    cell_text = "".join([n.text for n in t_nodes if n.text]).strip()
                    row_texts.append(cell_text)

                if len(row_texts) >= 3:
                    # check for RA reg number
                    reg = ""
                    name = ""
                    mobile = ""
                    email = ""
                    for i, txt in enumerate(row_texts):
                        t_upper = txt.upper()
                        if t_upper.startswith("RA") and len(t_upper) >= 10:
                            reg = txt
                            if i + 1 < len(row_texts):
                                name = row_texts[i + 1]
                            for rem in row_texts[i + 2:]:
                                if "@" in rem:
                                    email = rem
                                elif rem.replace(" ", "").isdigit() and len(rem.replace(" ", "")) >= 10:
                                    mobile = rem
                            break

                    if reg and name and not name.lower().startswith("name"):
                        patrons.append({
                            "barcode": f"ST{reg}",
                            "register_number": reg,
                            "name": name,
                            "patron_type": "student",
                            "year": year,
                            "section": section,
                            "mobile": mobile,
                            "email": email,
                            "parent_mobile": "",
                            "parent_email": "",
                        })
    except Exception as e:
        return [], [f"Error reading docx: {e}"]
    return patrons, []


def import_students_from_excel(filepath: str, year: str, section: str) -> tuple[list[dict], list[str]]:
    """
    Import students from class-wise Excel or docx files.
    """
    if str(filepath).lower().endswith(".docx"):
        return _import_students_from_docx(filepath, year, section)

    patrons = []
    errors = []

    def looks_like_reg(val):
        v = str(val).strip().upper()
        return bool(v and (v.startswith("RA") or (len(v) >= 10 and v[:2].isalpha() and v[2:].isdigit())))

    try:
        wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
    except Exception as e:
        return [], [f"Could not open file: {e}"]

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        if not hasattr(ws, 'max_row'):
            continue

        all_rows = list(ws.iter_rows(values_only=True))
        if not all_rows:
            continue

        header_idx = None
        header = None
        reg_col = None
        name_col = None
        mobile_col = None
        email_col = None
        pmob_col = None
        pemail_col = None

        for i, row in enumerate(all_rows):
            non_empty = [c for c in row if c is not None and str(c).strip()]
            if len(non_empty) < 3:
                continue

            row_lower = [str(c).strip().lower() for c in row if c is not None]
            has_reg = any(any(k in cell for k in ["reg. no", "reg no", "register", "registration", "reg number"]) for cell in row_lower)
            has_name = any(any(k in cell for k in ["name", "student name"]) for cell in row_lower)

            if has_reg and has_name:
                header_idx = i
                header = [str(c).strip() if c is not None else "" for c in row]
                break

        if header is not None and header_idx is not None:
            for c_idx, h in enumerate(header):
                hl = h.lower()
                if any(k in hl for k in ["reg. no", "reg no", "register", "registration", "reg number"]) and reg_col is None:
                    reg_col = c_idx
                elif any(k in hl for k in ["name", "student name"]) and name_col is None and "faculty" not in hl and "guide" not in hl:
                    name_col = c_idx
                elif any(k in hl for k in ["stu.mobile", "student mobile", "mobile no", "mobile number", "mobile"]) and "parent" not in hl and mobile_col is None:
                    mobile_col = c_idx
                elif any(k in hl for k in ["stu.email", "student email", "e-mail", "email", "mail"]) and "parent" not in hl and email_col is None:
                    email_col = c_idx
                elif "parent" in hl and any(k in hl for k in ["mobile", "phone"]) and pmob_col is None:
                    pmob_col = c_idx
                elif "parent" in hl and any(k in hl for k in ["email", "mail"]) and pemail_col is None:
                    pemail_col = c_idx

        # If header wasn't found by text, scan rows directly for data
        data_start = header_idx + 1 if header_idx is not None else 0
        for row in all_rows[data_start:]:
            if not row or not any(row):
                continue

            reg = ""
            name = ""
            mobile = ""
            email = ""
            pmob = ""
            pemail = ""

            if reg_col is not None and name_col is not None:
                if reg_col < len(row) and row[reg_col] is not None:
                    reg = _clean(row[reg_col])
                if name_col < len(row) and row[name_col] is not None:
                    name = _clean(row[name_col])
                if mobile_col is not None and mobile_col < len(row) and row[mobile_col] is not None:
                    mobile = _clean(row[mobile_col])
                if email_col is not None and email_col < len(row) and row[email_col] is not None:
                    email = _clean(row[email_col])
                if pmob_col is not None and pmob_col < len(row) and row[pmob_col] is not None:
                    pmob = _clean(row[pmob_col])
                if pemail_col is not None and pemail_col < len(row) and row[pemail_col] is not None:
                    pemail = _clean(row[pemail_col])
            else:
                # scan cells
                for idx, cell in enumerate(row):
                    if cell is not None and looks_like_reg(cell):
                        reg = _clean(cell)
                        if idx + 1 < len(row) and row[idx + 1]:
                            name = _clean(row[idx + 1])
                        break

            if not reg or not name:
                continue
            if not looks_like_reg(reg):
                continue
            if name.lower().startswith("name") or name.lower().startswith("student") or len(name) < 2:
                continue

            barcode = f"ST{reg}"
            patrons.append({
                "barcode": barcode,
                "register_number": reg,
                "name": name,
                "patron_type": "student",
                "year": year,
                "section": section,
                "mobile": mobile,
                "email": email,
                "parent_mobile": pmob,
                "parent_email": pemail,
            })

    wb.close()
    return patrons, errors


def import_staff_from_excel(filepath: str) -> tuple[list[dict], list[str]]:
    """
    Import teaching/non-teaching staff and research scholars from Excel file (.xls or .xlsx).
    Extracts name, ID, designation/category, guide (if RS), mobile, email.
    Returns (list_of_patrons, list_of_errors)
    """
    patrons = []
    errors = []
    is_xlsx = str(filepath).lower().endswith(".xlsx")

    sheets_data = []

    if is_xlsx:
        try:
            wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
            for sname in wb.sheetnames:
                ws = wb[sname]
                rows = list(ws.iter_rows(values_only=True))
                sheets_data.append((sname, rows))
            wb.close()
        except Exception as e:
            return [], [f"Could not open XLSX file: {e}"]
    else:
        try:
            import xlrd
            wb = xlrd.open_workbook(filepath)
            for sheet_idx in range(wb.nsheets):
                ws = wb.sheet_by_index(sheet_idx)
                rows = []
                for rx in range(ws.nrows):
                    rows.append([ws.cell_value(rx, cx) for cx in range(ws.ncols)])
                sheets_data.append((ws.name.strip(), rows))
        except Exception as e:
            return [], [f"Could not open XLS file: {e}"]

    for sheet_idx, (sheet_name, all_rows) in enumerate(sheets_data):
        if not all_rows:
            continue

        header_row_idx = None
        id_col = None
        name_col = None
        desig_col = None
        guide_col = None
        mobile_col = None
        email_col = None

        for rx, row in enumerate(all_rows[:15]):
            row_vals = [str(c or "").strip().lower() for c in row]
            has_id = any(any(k in c for k in ["id. no", "id no", "register number", "emp id", "staff id", "sl. no", "reg no"]) for c in row_vals)
            has_name = any("name" in c for c in row_vals)
            if has_id and has_name:
                header_row_idx = rx
                for cx, c_val in enumerate(row_vals):
                    if any(k in c_val for k in ["id. no", "id no", "register number", "emp id", "reg no", "staff id"]):
                        id_col = cx
                    elif "name" in c_val and "guide" not in c_val:
                        name_col = cx
                    elif any(k in c_val for k in ["designation", "desg", "post", "category"]):
                        desig_col = cx
                    elif "guide" in c_val:
                        guide_col = cx
                    elif any(k in c_val for k in ["mobile", "contact"]):
                        mobile_col = cx
                    elif any(k in c_val for k in ["mail", "email"]):
                        email_col = cx
                break

        if header_row_idx is None or name_col is None:
            continue

        for rx in range(header_row_idx + 1, len(all_rows)):
            row = all_rows[rx]
            if not row or not any(row):
                continue

            raw_name = _clean(row[name_col]) if name_col < len(row) else ""
            if not raw_name or raw_name.lower().startswith("name") or len(raw_name) < 2:
                continue

            raw_id = _clean(row[id_col]) if id_col is not None and id_col < len(row) else ""
            if not raw_id:
                raw_id = f"STAFF{sheet_idx+1}_{rx}"

            raw_desig = _clean(row[desig_col]) if desig_col is not None and desig_col < len(row) else ""
            raw_guide = _clean(row[guide_col]) if guide_col is not None and guide_col < len(row) else ""
            full_desig = raw_desig
            if raw_guide and raw_desig:
                full_desig = f"{raw_desig} (Guide: {raw_guide})"
            elif raw_guide:
                full_desig = f"Guide: {raw_guide}"

            raw_mobile = _clean(row[mobile_col]) if mobile_col is not None and mobile_col < len(row) else ""
            raw_email = _clean(row[email_col]) if email_col is not None and email_col < len(row) else ""

            # Standardize barcode
            clean_id = raw_id.replace(" ", "").replace("/", "").replace(".", "")
            barcode = f"TC{clean_id}"

            p_type = "teacher" if "rs" not in sheet_name.lower() else "student"
            patrons.append({
                "barcode": barcode,
                "register_number": raw_id,
                "name": raw_name,
                "patron_type": p_type,
                "designation": full_desig,
                "year": "RS" if p_type == "student" else "",
                "section": full_desig if full_desig else sheet_name,
                "mobile": raw_mobile,
                "email": raw_email,
                "parent_mobile": "",
                "parent_email": "",
            })

    return patrons, errors
